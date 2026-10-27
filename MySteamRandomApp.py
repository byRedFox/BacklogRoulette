"""My Steam Random — рандомайзер «Во что сыграть дальше?» для библиотеки Steam."""

import io
import json
import os
import random
import re
import sys
import tkinter as tk
from tkinter import filedialog, messagebox
from urllib.parse import quote

import requests
from PIL import Image, ImageTk

# --- Пути -----------------------------------------------------------------
# Данные лежат рядом со скриптом (или рядом с .exe после сборки PyInstaller),
# поэтому программа работает из любой папки на любом компьютере.
if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Ресурсы, вшитые в .exe, распаковываются во временную папку _MEIPASS
RESOURCE_DIR = getattr(sys, "_MEIPASS", BASE_DIR)

GAMES_FILE = os.path.join(BASE_DIR, "games.json")
LEGACY_EXCEL_FILE = os.path.join(BASE_DIR, "SteamGames.xlsx")
COVERS_DIR = os.path.join(BASE_DIR, "covers")
API_KEY_FILE = os.path.join(BASE_DIR, "api_key.txt")  # ключ SteamGridDB
SETTINGS_FILE = os.path.join(BASE_DIR, "settings.json")  # ключ Steam и SteamID
ICON_FILE = os.path.join(RESOURCE_DIR, "1.ico")

# --- Настройки приложения ---------------------------------------------------
WINDOW_SIZE = "460x760"
FONT = ("Comic Sans MS", 12)
TITLE_FONT = ("Comic Sans MS", 16, "bold")
SMALL_FONT = ("Segoe UI", 9)
COVER_SIZE = (300, 450)
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg")

STEAMGRIDDB_API_URL = "https://www.steamgriddb.com/api/v2"
STEAM_OWNED_GAMES_URL = "https://api.steampowered.com/IPlayerService/GetOwnedGames/v1/"
STEAM_COVER_URL = "https://cdn.cloudflare.steamstatic.com/steam/apps/{appid}/library_600x900.jpg"
REQUEST_TIMEOUT = 10  # секунд


# --- Работа с файлами -------------------------------------------------------
def read_text(path):
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return f.read().strip()
    return ""


def write_text(path, text):
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


def save_json(path, data):
    # Пишем во временный файл и подменяем, чтобы не испортить список при сбое
    tmp_path = path + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    os.replace(tmp_path, path)


def load_settings():
    try:
        with open(SETTINGS_FILE, encoding="utf-8") as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return {}


# Очистка названия от символов, недопустимых в имени файла (и значков ™ ®)
def clean_name(text):
    text = re.sub(r'[<>:"/\\|?*™®©]', "", str(text))
    return re.sub(r"\s+", " ", text).strip()


# --- Список игр -------------------------------------------------------------
def read_legacy_excel():
    """Читает названия из старого SteamGames.xlsx (столбец «Игры»)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        messagebox.showwarning(
            "Перенос списка",
            "Найден старый SteamGames.xlsx, но для переноса нужен пакет openpyxl.\n\n"
            "Выполните: pip install openpyxl\nи запустите программу снова.",
        )
        return None

    workbook = load_workbook(LEGACY_EXCEL_FILE, read_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    header = next(rows, ())
    column = header.index("Игры") if "Игры" in header else 0
    names = [row[column] for row in rows if len(row) > column and row[column]]
    workbook.close()
    return names


class GameLibrary:
    """Список игр в games.json: [{"name": "...", "appid": 123 или null}, ...]"""

    def __init__(self, path):
        self.path = path
        self.games = []

    def load(self):
        if not os.path.exists(self.path):
            self._migrate_from_excel()
            return
        try:
            with open(self.path, encoding="utf-8") as f:
                data = json.load(f)
        except json.JSONDecodeError as e:
            # Не перезаписываем повреждённый файл — откладываем его в сторону
            broken_path = self.path + ".broken"
            os.replace(self.path, broken_path)
            messagebox.showerror(
                "Ошибка",
                f"Файл games.json повреждён ({e}).\n"
                f"Он сохранён как {os.path.basename(broken_path)}, список начат заново.",
            )
            data = []
        self.games = [g for g in data if isinstance(g, dict) and g.get("name")]

    def _migrate_from_excel(self):
        if not os.path.exists(LEGACY_EXCEL_FILE):
            return
        names = read_legacy_excel()
        if not names:
            return
        for name in names:
            self.add(name)
        self.save()
        messagebox.showinfo(
            "Перенос списка",
            f"Из SteamGames.xlsx перенесено игр: {len(self.games)}.\n"
            "Теперь список хранится в games.json.",
        )

    def save(self):
        save_json(self.path, self.games)

    def find(self, name):
        key = name.casefold()
        for game in self.games:
            if game["name"].casefold() == key:
                return game
        return None

    def add(self, name, appid=None):
        """Возвращает "added", "updated" (дописан appid) или None (пропущено)."""
        name = clean_name(name)
        if not name:
            return None
        existing = self.find(name)
        if existing:
            if appid and not existing.get("appid"):
                existing["appid"] = appid
                return "updated"
            return None
        self.games.append({"name": name, "appid": appid})
        return "added"

    def remove(self, game):
        self.games.remove(game)


# --- Обложки ----------------------------------------------------------------
def find_cover(name):
    base = clean_name(name)
    for ext in IMAGE_EXTENSIONS:
        path = os.path.join(COVERS_DIR, base + ext)
        if os.path.exists(path):
            return path
    return None


def delete_cover(name):
    base = clean_name(name)
    for ext in IMAGE_EXTENSIONS:
        path = os.path.join(COVERS_DIR, base + ext)
        if os.path.exists(path):
            os.remove(path)


def save_cover_bytes(name, content):
    image = Image.open(io.BytesIO(content)).convert("RGB")
    image = image.resize(COVER_SIZE, Image.Resampling.LANCZOS)
    os.makedirs(COVERS_DIR, exist_ok=True)
    delete_cover(name)  # иначе старая .png найдётся раньше новой .jpg
    path = os.path.join(COVERS_DIR, clean_name(name) + ".jpg")
    image.save(path, quality=92)
    return path


def fetch(url, **kwargs):
    response = requests.get(url, timeout=REQUEST_TIMEOUT, **kwargs)
    response.raise_for_status()
    return response


def download_from_steamgriddb(name, api_key):
    headers = {"Authorization": f"Bearer {api_key}"}
    search_url = f"{STEAMGRIDDB_API_URL}/search/autocomplete/{quote(name, safe='')}"
    found = fetch(search_url, headers=headers).json().get("data", [])
    if not found:
        return None

    grids_url = f"{STEAMGRIDDB_API_URL}/grids/game/{found[0]['id']}"
    # Сначала просим вертикальные обложки 600x900, затем любые
    grids = fetch(grids_url, headers=headers, params={"dimensions": "600x900"}).json().get("data", [])
    if not grids:
        grids = fetch(grids_url, headers=headers).json().get("data", [])
    if not grids:
        return None
    return fetch(grids[0]["url"]).content


def download_cover(game, steamgriddb_key):
    """Возвращает (путь к обложке или None, список ошибок)."""
    errors = []

    # 1. Официальная обложка Steam — не требует ключа, нужен только appid
    if game.get("appid"):
        try:
            content = fetch(STEAM_COVER_URL.format(appid=game["appid"])).content
            return save_cover_bytes(game["name"], content), errors
        except (requests.RequestException, OSError) as e:
            errors.append(f"Steam: {e}")

    # 2. SteamGridDB — поиск по названию
    if not steamgriddb_key:
        errors.append("SteamGridDB: не задан API-ключ")
        return None, errors
    try:
        content = download_from_steamgriddb(game["name"], steamgriddb_key)
        if content:
            return save_cover_bytes(game["name"], content), errors
        errors.append("SteamGridDB: игра не найдена")
    except (requests.RequestException, OSError, ValueError, KeyError) as e:
        errors.append(f"SteamGridDB: {e}")
    return None, errors


# --- Steam ------------------------------------------------------------------
def fetch_owned_games(api_key, steam_id):
    params = {
        "key": api_key,
        "steamid": steam_id,
        "include_appinfo": 1,
        "include_played_free_games": 1,
        "format": "json",
    }
    data = fetch(STEAM_OWNED_GAMES_URL, params=params).json()
    return data.get("response", {}).get("games", [])


# --- Интерфейс --------------------------------------------------------------
def center_window(window):
    window.update_idletasks()
    width = window.winfo_width()
    height = window.winfo_height()
    x = (window.winfo_screenwidth() // 2) - (width // 2)
    y = (window.winfo_screenheight() // 2) - (height // 2)
    window.geometry(f"{width}x{height}+{x}+{y}")


def set_window_icon(window):
    if os.path.exists(ICON_FILE):
        try:
            window.iconbitmap(ICON_FILE)
        except tk.TclError:
            pass


def make_dialog(root, title):
    dialog = tk.Toplevel(root)
    dialog.title(title)
    dialog.resizable(False, False)
    dialog.transient(root)  # связываем с главным окном
    dialog.grab_set()  # делаем окно модальным
    set_window_icon(dialog)
    return dialog


class App:
    def __init__(self, root):
        self.root = root
        self.library = GameLibrary(GAMES_FILE)
        self.current = None
        self.cover_photo = None  # держим ссылку, иначе картинку съест сборщик мусора

        root.title("My Steam Random")
        root.geometry(WINDOW_SIZE)
        root.resizable(False, False)
        set_window_icon(root)

        self.title_label = tk.Label(
            root, text="Во что сыграть дальше?", font=TITLE_FONT, wraplength=420
        )
        self.title_label.pack(pady=(15, 10))

        self.blank_cover = tk.PhotoImage(width=COVER_SIZE[0], height=COVER_SIZE[1])
        self.cover_label = tk.Label(
            root,
            image=self.blank_cover,
            text="Нажмите «Выбрать игру»",
            compound="center",
            font=FONT,
            relief="groove",
        )
        self.cover_label.pack()

        play_row = tk.Frame(root)
        play_row.pack(pady=(15, 5))
        self.pick_button = tk.Button(
            play_row, text="Выбрать игру", font=FONT, width=16, command=self.pick_game
        )
        self.pick_button.pack(side="left", padx=5)
        self.done_button = tk.Button(
            play_row,
            text="Пройдена — убрать",
            font=FONT,
            width=16,
            state="disabled",
            command=self.remove_current,
        )
        self.done_button.pack(side="left", padx=5)

        manage_row = tk.Frame(root)
        manage_row.pack(pady=5)
        tk.Button(manage_row, text="Добавить", command=self.add_games_dialog).pack(side="left", padx=3)
        tk.Button(manage_row, text="Импорт из Steam", command=self.import_from_steam).pack(side="left", padx=3)
        self.cover_button = tk.Button(
            manage_row, text="Своя обложка", state="disabled", command=self.choose_custom_cover
        )
        self.cover_button.pack(side="left", padx=3)
        tk.Button(manage_row, text="Настройки", command=self.settings_dialog).pack(side="left", padx=3)

        self.status_label = tk.Label(root, text="", font=SMALL_FONT, fg="gray", wraplength=420)
        self.status_label.pack(pady=(8, 0))
        self.count_label = tk.Label(root, text="", font=SMALL_FONT)
        self.count_label.pack(side="bottom", pady=8)

        root.bind("<Return>", lambda event: self.pick_game())

        center_window(root)
        self.library.load()
        self.update_count()

    # --- Вспомогательное ---
    def set_status(self, text):
        self.status_label.config(text=text)

    def update_count(self):
        self.count_label.config(text=f"Игр в списке: {len(self.library.games)}")

    def save_library(self):
        try:
            self.library.save()
            return True
        except OSError as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить список игр: {e}")
            return False

    def busy(self, is_busy):
        self.root.config(cursor="watch" if is_busy else "")
        self.root.update_idletasks()

    def show_cover(self, path):
        if path:
            try:
                image = Image.open(path).resize(COVER_SIZE, Image.Resampling.LANCZOS)
                self.cover_photo = ImageTk.PhotoImage(image)
                self.cover_label.config(image=self.cover_photo, text="")
                return
            except OSError as e:
                self.set_status(f"Не удалось открыть обложку: {e}")
        self.cover_label.config(image=self.blank_cover, text="Нет обложки")

    def reset_selection(self):
        self.current = None
        self.title_label.config(text="Во что сыграть дальше?")
        self.pick_button.config(text="Выбрать игру")
        self.done_button.config(state="disabled")
        self.cover_button.config(state="disabled")
        self.cover_label.config(image=self.blank_cover, text="Нажмите «Выбрать игру»")

    # --- Выбор игры ---
    def pick_game(self):
        games = self.library.games
        if not games:
            if messagebox.askyesno("Список пуст", "В списке нет игр. Добавить?"):
                self.add_games_dialog()
            return
        # Не показываем ту же игру два раза подряд
        candidates = [g for g in games if g is not self.current] or games
        self.show_game(random.choice(candidates))

    def show_game(self, game):
        self.current = game
        self.title_label.config(text=game["name"])
        self.pick_button.config(text="Другая игра")
        self.done_button.config(state="normal")
        self.cover_button.config(state="normal")
        self.set_status("")

        path = find_cover(game["name"])
        if not path:
            self.set_status("Ищу обложку…")
            self.busy(True)
            try:
                path, errors = download_cover(game, read_text(API_KEY_FILE))
            finally:
                self.busy(False)
            if path:
                self.set_status("")
            else:
                self.set_status(
                    "Обложка не найдена (" + "; ".join(errors) + "). "
                    "Можно выбрать свою кнопкой «Своя обложка»."
                )
        self.show_cover(path)

    def remove_current(self):
        if not self.current:
            return
        name = self.current["name"]
        if not messagebox.askyesno("Убрать игру", f"Убрать «{name}» из списка?"):
            return
        self.library.remove(self.current)
        if not self.save_library():
            return
        try:
            delete_cover(name)
        except OSError:
            pass  # обложка не критична
        self.reset_selection()
        self.update_count()
        self.set_status(f"«{name}» убрана из списка. Отличная работа!")

    def choose_custom_cover(self):
        if not self.current:
            return
        path = filedialog.askopenfilename(
            title="Выберите обложку для игры",
            filetypes=[("Изображения", "*.png *.jpg *.jpeg *.webp")],
        )
        if not path:
            return
        try:
            with open(path, "rb") as f:
                cover_path = save_cover_bytes(self.current["name"], f.read())
        except OSError as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить обложку: {e}")
            return
        self.set_status("Обложка сохранена.")
        self.show_cover(cover_path)

    # --- Добавление игр ---
    def add_games_dialog(self):
        dialog = make_dialog(self.root, "Добавить игры")

        tk.Label(dialog, text="По одной игре в строке (или через «|»):", font=FONT).pack(
            padx=15, pady=(15, 5)
        )
        text = tk.Text(dialog, font=FONT, width=40, height=8)
        text.pack(padx=15)
        text.focus_set()

        def paste():
            try:
                text.insert(tk.INSERT, self.root.clipboard_get())
            except tk.TclError:
                pass  # в буфере нет текста

        def submit():
            names = re.split(r"[|\n]", text.get("1.0", "end"))
            added = [clean_name(n) for n in names if self.library.add(n) == "added"]
            if not added:
                messagebox.showwarning(
                    "Добавить игры", "Нечего добавлять: поле пустое или игры уже в списке.", parent=dialog
                )
                return
            if not self.save_library():
                return
            dialog.destroy()
            self.update_count()
            messagebox.showinfo(
                "Готово",
                f"Добавлено игр: {len(added)}\n\n" + "\n".join(added[:15]) + ("\n…" if len(added) > 15 else ""),
            )

        buttons = tk.Frame(dialog)
        buttons.pack(pady=10)
        tk.Button(buttons, text="Вставить из буфера", font=FONT, command=paste).pack(side="left", padx=5)
        tk.Button(buttons, text="Добавить", font=FONT, command=submit).pack(side="left", padx=5)
        center_window(dialog)

    def import_from_steam(self):
        settings = load_settings()
        api_key = settings.get("steam_api_key", "")
        steam_id = settings.get("steam_id", "")
        if not api_key or not steam_id:
            messagebox.showinfo(
                "Импорт из Steam", "Сначала укажите Steam Web API key и SteamID64 в настройках."
            )
            self.settings_dialog()
            return

        only_unplayed = messagebox.askyesnocancel(
            "Импорт из Steam",
            "Импортировать только игры, в которые вы ещё не играли?\n\n"
            "Да — только неначатые\nНет — всю библиотеку",
        )
        if only_unplayed is None:
            return

        self.set_status("Загружаю библиотеку Steam…")
        self.busy(True)
        try:
            owned = fetch_owned_games(api_key, steam_id)
        except (requests.RequestException, ValueError) as e:
            self.set_status("")
            messagebox.showerror("Ошибка", f"Не удалось получить библиотеку Steam:\n{e}")
            return
        finally:
            self.busy(False)

        if not owned:
            self.set_status("")
            messagebox.showwarning(
                "Импорт из Steam",
                "Steam не вернул ни одной игры.\n\nПроверьте SteamID64 и что в настройках "
                "приватности профиля Steam «Игровая информация» открыта для всех.",
            )
            return

        added = updated = 0
        for game in owned:
            if only_unplayed and game.get("playtime_forever", 0) > 0:
                continue
            result = self.library.add(game.get("name", ""), game.get("appid"))
            added += result == "added"
            updated += result == "updated"

        if (added or updated) and not self.save_library():
            return
        self.update_count()
        self.set_status("")
        messagebox.showinfo(
            "Импорт из Steam",
            f"Игр в библиотеке Steam: {len(owned)}\n"
            f"Добавлено новых: {added}\n"
            f"Уже были в списке (дописан Steam ID для обложек): {updated}",
        )

    # --- Настройки ---
    def settings_dialog(self):
        settings = load_settings()
        dialog = make_dialog(self.root, "Настройки")

        fields = [
            ("SteamGridDB API key", read_text(API_KEY_FILE),
             "steamgriddb.com → Preferences → API (обложки)"),
            ("Steam Web API key", settings.get("steam_api_key", ""),
             "steamcommunity.com/dev/apikey (импорт библиотеки)"),
            ("SteamID64", settings.get("steam_id", ""),
             "17 цифр, например 7656119…; можно узнать на steamid.io"),
        ]
        entries = []
        for row, (label, value, hint) in enumerate(fields):
            tk.Label(dialog, text=label, font=FONT).grid(row=row * 2, column=0, sticky="w", padx=15, pady=(10, 0))
            entry = tk.Entry(dialog, font=FONT, width=38)
            entry.insert(0, value)
            entry.grid(row=row * 2, column=1, padx=(0, 15), pady=(10, 0))
            tk.Label(dialog, text=hint, font=SMALL_FONT, fg="gray").grid(
                row=row * 2 + 1, column=1, sticky="w", padx=(0, 15)
            )
            entries.append(entry)

        def save():
            griddb_key, steam_key, steam_id = (e.get().strip() for e in entries)
            if steam_id and not re.fullmatch(r"\d{17}", steam_id):
                messagebox.showwarning("Настройки", "SteamID64 должен состоять из 17 цифр.", parent=dialog)
                return
            try:
                write_text(API_KEY_FILE, griddb_key)
                save_json(SETTINGS_FILE, {"steam_api_key": steam_key, "steam_id": steam_id})
            except OSError as e:
                messagebox.showerror("Ошибка", f"Не удалось сохранить настройки: {e}", parent=dialog)
                return
            dialog.destroy()
            self.set_status("Настройки сохранены.")

        tk.Button(dialog, text="Сохранить", font=FONT, command=save).grid(
            row=len(fields) * 2, column=0, columnspan=2, pady=15
        )
        center_window(dialog)


def main():
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
